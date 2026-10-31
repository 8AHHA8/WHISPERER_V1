import os
import time
import tempfile
import yt_dlp
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from faster_whisper import WhisperModel
from assets.config import MODEL_NAME, DEVICE, COMPUTE_TYPE
from core.core_docx import make_docx, save_docx_safely

def download_youtube_audio(url):
    temp_dir = tempfile.gettempdir()
    outtmpl = os.path.join(temp_dir, "%(title)s.%(ext)s")
    ydl_opts = {'format': 'bestaudio/best', 'outtmpl': outtmpl, 'noplaylist': False, 'quiet': True, 'no_warnings': True}
    downloaded_files = []
    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(url, download=True)
        if 'entries' in info:
            for entry in info['entries']:
                if entry:
                    file_path = ydl.prepare_filename(entry)
                    if os.path.exists(file_path):
                        downloaded_files.append((file_path, entry.get('title', 'Unknown_title')))
        else:
            file_path = ydl.prepare_filename(info)
            if os.path.exists(file_path):
                downloaded_files.append((file_path, info.get('title', 'Unknown_title')))
        if not downloaded_files:
            raise Exception(f"Failed to download audio from: {url}")
        return downloaded_files

def transcribe_file(model, audio_path, index, total, destination, events):
    is_url = audio_path.startswith("http")
    temp_audio_files = []
    
    if is_url:
        events.put(("log", f"Downloading audio from: {audio_path}"))
        try:
            temp_audio_files = download_youtube_audio(audio_path)
            events.put(("log", f"Temporarily downloaded {len(temp_audio_files)} file(s)."))
        except Exception as e:
            events.put(("error", f"Download error: {e}"))
            return
    else:
        temp_audio_files = [(audio_path, os.path.splitext(os.path.basename(audio_path))[0])]

    for file_idx, (current_audio, video_title) in enumerate(temp_audio_files, start=1):
        safe_name = "".join([c for c in video_title if c not in r'<>:"/\|?*'])
        docx_path = os.path.join(destination, safe_name + ".docx")
        display_name = f"{video_title}" if len(temp_audio_files) == 1 else f"{video_title} ({file_idx}/{len(temp_audio_files)})"
        
        events.put(("status", f"Processing {index}/{total}: {display_name}"))
        events.put(("log", f"Started: {display_name}"))

        # core_audio.py (w model.transcribe)
        segments, info = model.transcribe(
            current_audio, 
            language="en", 
            word_timestamps=True, 
            vad_filter=True,             
            beam_size=10,                
            patience=1.0,                
            condition_on_previous_text=False, 
            temperature=0.0,             
            repetition_penalty=1.2,      
            no_repeat_ngram_size=3,      
            hallucination_silence_threshold=1.0, 
            compression_ratio_threshold=2.4,
            suppress_tokens=[-1],
        )
        duration = info.duration or 0
        next_minute = 1
        last_percent = -1
        last_save = time.monotonic()
        word_count = 0

        doc = make_docx()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for segment in segments:
            if not segment.words: continue
            for word in segment.words:
                word_count += 1
                while word.start >= next_minute * 60:
                    marker_paragraph = doc.add_paragraph()
                    marker = marker_paragraph.add_run(f"[{next_minute} min]")
                    marker.bold = True
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    next_minute += 1

                run = paragraph.add_run(word.word)
                percent = min(100, int(word.end / duration * 100)) if duration else 0

                if percent != last_percent:
                    base_progress = ((index - 1) + (file_idx - 1) / len(temp_audio_files)) / total * 100
                    file_progress = (percent / 100) / len(temp_audio_files) / total * 100
                    events.put(("progress", base_progress + file_progress, f"{index}/{total}: {display_name} — {percent}%"))
                    last_percent = percent

                if time.monotonic() - last_save >= 10:
                    if not save_docx_safely(doc, docx_path):
                        events.put(("warning", f"{display_name}: DOCX file is locked."))
                    last_save = time.monotonic()

        if word_count == 0: events.put(("warning", f"{display_name}: No words recognized."))
        if save_docx_safely(doc, docx_path): events.put(("log", f"Completed: {display_name}"))
        
        if is_url and current_audio and os.path.exists(current_audio):
            try: os.remove(current_audio)
            except: pass

    events.put(("progress", index / total * 100, f"Finished file {index}/{total}"))

def worker(files_to_process, destination, events):
    try:
        events.put(("status", "Loading Whisper model..."))
        model = WhisperModel(MODEL_NAME, device=DEVICE, compute_type=COMPUTE_TYPE)
        events.put(("log", f"Model loaded."))
        total = len(files_to_process)

        for index, audio_path in enumerate(files_to_process, start=1):
            try: transcribe_file(model, audio_path, index, total, destination, events)
            except Exception as error: events.put(("error", f"File error: {error}"))

        events.put(("progress", 100, "Processing finished."))
        events.put(("finished", "All files processed successfully."))
    except Exception as error:
        events.put(("error", f"Failed to load model: {error}"))
        events.put(("finished", "Processing interrupted."))
    finally:
        events.put(("worker_stopped",))